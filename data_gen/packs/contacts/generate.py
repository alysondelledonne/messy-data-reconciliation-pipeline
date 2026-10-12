"""Contacts pack: 300 true people, each mentioned across emails, Word docs, notes,
JSONL, logs and a CRM export.

Planted mess:
- nicknames, "Last, First" and uppercase names, partial names such as "Dana W.";
- company alias spellings, and lookalike companies that must *not* be merged;
- namesakes: different people who share a name;
- phone numbers in many formats, some with transposed digits;
- email typos and inconsistent case;
- stale job titles, phone numbers or employers from before someone changed jobs;
- quoted reply chains that mention a *different* person;
- vCard attachments, multipart HTML email, inconsistent JSON key names, malformed
  JSONL lines, log noise, and duplicate CRM rows and form submissions.
"""

import html
import json
import random
import unicodedata
from collections.abc import Mapping
from dataclasses import dataclass, field
from datetime import UTC, date, datetime, timedelta
from email.message import EmailMessage
from email.utils import format_datetime
from typing import Any

from docx import Document

from data_gen import GENERATOR_VERSION
from data_gen.corpus import PackOutput, sub_rng
from data_gen.ground_truth import (
    Entity,
    Format,
    GroundTruth,
    Occurrence,
    TrueRecord,
    occurrence_id,
)
from data_gen.mess import (
    NICKNAMES,
    PHONE_STYLES,
    format_phone,
    random_org_variant,
    transpose_digits,
    typo,
)
from data_gen.packs.contacts.vocab import (
    BODY_LINES,
    CITIES,
    COMPANIES,
    COUNTRY_SPELLINGS,
    FIRST_NAMES,
    JOB_TITLES,
    LAST_NAMES,
    SUBJECTS,
    THINGS,
    City,
    Company,
)
from data_gen.writers.files import csv_bytes, docx_bytes, eml_bytes
from pipeline.domains.contacts.schema import Contact

PACK = "contacts"


@dataclass(frozen=True)
class ContactsConfig:
    n_records: int = 300
    n_namesake_pairs: int = 3
    history_share: float = 0.18  # people with a stale title/phone/employer somewhere
    n_eml: int = 80
    n_eml_senders: int = 55
    reply_chain_share: float = 0.25
    vcard_share: float = 0.15
    html_share: float = 0.3
    n_docx: int = 20
    n_notes: int = 40
    n_jsonl_lines: int = 120
    n_jsonl_noise: int = 6
    n_jsonl_dups: int = 8
    n_log_lines: int = 100
    log_contact_share: float = 0.45
    crm_share: float = 0.65
    n_crm_dups: int = 10
    start: date = date(2023, 1, 1)
    end: date = date(2025, 12, 31)


@dataclass
class _Person:
    record_id: str
    first: str
    last: str
    company: Company
    title: str
    phone: str
    email: str
    city: City
    old_title: str | None = None
    old_phone: str | None = None
    old_company: Company | None = None
    old_email: str | None = None
    change_date: date | None = None  # sources dated before this show the old values
    occurrences: list[Occurrence] = field(default_factory=list)

    def truth(self) -> Contact:
        return Contact(
            first_name=self.first,
            last_name=self.last,
            email=self.email,
            phone=self.phone,
            company=self.company.name,
            job_title=self.title,
            city=self.city.name,
            country=self.city.country,
        )

    @property
    def has_history(self) -> bool:
        return bool(self.old_title or self.old_phone or self.old_company)


def _ascii(s: str) -> str:
    norm = unicodedata.normalize("NFKD", s)
    return "".join(ch for ch in norm if ch.isascii() and ch.isalpha()).lower()


def _email(first: str, last: str, domain: str, pattern: int) -> str:
    f, la = _ascii(first), _ascii(last)
    local = [f"{f}.{la}", f"{f[0]}{la}", f"{f}{la[0]}", f"{f}_{la}"][pattern]
    return f"{local}@{domain}"


class _PhonePool:
    def __init__(self, rng: random.Random) -> None:
        self.rng = rng
        self.free: dict[str, list[str]] = {}

    def draw(self, city: City) -> str:
        if city.name not in self.free:
            if city.country == "GB":
                nums = [f"{city.dial}{n:04d}" for n in range(1000)]  # 0000..0999
            else:
                nums = [f"{city.dial}01{n:02d}" for n in range(100)]  # 0100..0199
            self.rng.shuffle(nums)
            self.free[city.name] = nums
        return self.free[city.name].pop()


def _make_people(cfg: ContactsConfig, rng: random.Random) -> list[_Person]:
    names = [(f, la) for f in FIRST_NAMES for la in LAST_NAMES]
    picked = rng.sample(names, cfg.n_records - cfg.n_namesake_pairs)
    picked += rng.sample(picked, cfg.n_namesake_pairs)  # namesakes: same name, other person
    rng.shuffle(picked)
    weights = [c.weight for c in COMPANIES]
    patterns = {c.company_id: rng.randrange(4) for c in COMPANIES}
    phones = _PhonePool(rng)
    people: list[_Person] = []
    used_emails: set[str] = set()
    for i, (first, last) in enumerate(picked, 1):
        company = rng.choices(COMPANIES, weights)[0]
        city = CITIES[company.city]
        email = _email(first, last, company.domain, patterns[company.company_id])
        while email in used_emails:  # namesakes at the same company
            local, dom = email.split("@")
            email = f"{local}{rng.randint(2, 9)}@{dom}"
        used_emails.add(email)
        p = _Person(f"con-{i:04d}", first, last, company, rng.choice(JOB_TITLES),
                    phones.draw(city), email, city)  # fmt: skip
        if rng.random() < cfg.history_share:
            kind = rng.choices(["title", "phone", "company"], [5, 3, 2])[0]
            if kind == "title":
                p.old_title = rng.choice([t for t in JOB_TITLES if t != p.title])
            elif kind == "phone":
                p.old_phone = phones.draw(city)
            else:
                old = rng.choice([c for c in COMPANIES if c.company_id != company.company_id])
                p.old_company = old
                p.old_email = _email(first, last, old.domain, patterns[old.company_id])
                p.old_title = rng.choice(JOB_TITLES)
            span = (cfg.end - cfg.start).days
            p.change_date = cfg.start + timedelta(days=rng.randint(120, span - 120))
        people.append(p)
    return people


# --- one rendering of a person, with planted mess --------------------------------------


@dataclass
class _View:
    first: str | None
    last: str | None
    email: str | None
    phone: str | None  # E.164 of the number as written (may be stale or wrong)
    phone_text: str | None
    company: str | None
    title: str | None
    city: str | None
    country_text: str | None
    country: str | None
    defects: list[str]

    def name(self, style: str = "normal") -> str:
        f, la = self.first or "", self.last or ""
        match style:
            case "last_first":
                return f"{la}, {f}"
            case "upper_last_first":
                return f"{la.upper()}, {f.upper()}"
            case "partial":
                return f"{f} {la[:1]}."
        return f"{f} {la}".strip()

    def expected(self, *, partial_name: bool = False) -> dict[str, Any]:
        exp: dict[str, Any] = {
            "first_name": self.first,
            "last_name": None if partial_name else self.last,
            "email": self.email,
            "phone": self.phone,
            "company": self.company,
            "job_title": self.title,
            "city": self.city,
            "country": self.country,
        }
        return {k: v for k, v in exp.items() if v is not None}


def _view(
    p: _Person,
    rng: random.Random,
    *,
    fields: tuple[str, ...] = ("email", "phone", "company", "title", "city", "country"),
    p_nick: float = 0.1,
    p_alias: float = 0.25,
    when: date | None = None,
    p_drop: float = 0.0,
    p_typo: float = 0.03,
    canonical: bool = False,
) -> _View:
    defects: list[str] = []
    first = p.first
    if not canonical and first in NICKNAMES and rng.random() < p_nick:
        first, defects = NICKNAMES[first], [*defects, "nickname"]
    # Values change on the person's change date: a source dated before it shows the old ones.
    stale = (
        not canonical and p.change_date is not None and when is not None and when < p.change_date
    )
    company_obj = p.old_company if stale and p.old_company else p.company
    title = p.old_title if stale and p.old_title else p.title
    phone = p.old_phone if stale and p.old_phone else p.phone
    email = p.old_email if stale and p.old_email else p.email
    if stale:
        if p.old_company:
            defects.append("stale_company")
        elif p.old_title:
            defects.append("stale_title")
        if p.old_phone:
            defects.append("stale_phone")

    def keep(name: str) -> bool:
        return name in fields and (canonical or rng.random() >= p_drop)

    company = company_obj.name
    if not canonical and rng.random() < p_alias:
        company = random_org_variant(company, rng)
        if company != company_obj.name:
            defects.append("company_alias")
    phone_text = format_phone(phone, rng.choice(PHONE_STYLES))
    if not canonical and rng.random() < p_typo:
        phone_text = transpose_digits(phone_text, rng)
        defects.append("phone_transposed")
    if not canonical and rng.random() < p_typo:
        local, dom = email.split("@")
        email = f"{local}@{typo(dom, rng)}" if rng.random() < 0.5 else f"{typo(local, rng)}@{dom}"
        defects.append("email_typo")
    if not canonical and rng.random() < 0.08:
        email = email[0].upper() + email[1:]
    country_text = rng.choice(COUNTRY_SPELLINGS[p.city.country])
    # Decide once per field: the written text and `expected` must agree.
    kept = {f: keep(f) for f in ("email", "phone", "company", "title", "city", "country")}
    return _View(
        first=first,
        last=p.last,
        email=email if kept["email"] else None,
        phone=_e164(phone_text) if kept["phone"] else None,
        phone_text=phone_text if kept["phone"] else None,
        company=company if kept["company"] else None,
        title=title if kept["title"] else None,
        city=p.city.name if kept["city"] else None,
        country_text=country_text if kept["country"] else None,
        country=p.city.country if kept["country"] else None,
        defects=defects,
    )


def _e164(text: str) -> str:
    """E.164 form of a human-written number (extension dropped)."""
    text = text.split(" x")[0]
    digits = "".join(ch for ch in text if ch.isdigit())
    if text.startswith("+"):
        return "+" + digits
    if digits.startswith("0"):  # UK national format
        return "+44" + digits[1:]
    return "+1" + digits


def _occ(
    p: _Person,
    source: str,
    fmt: Format,
    loc: Mapping[str, int | str],
    expected: dict[str, Any],
    defects: list[str],
    duplicate_of: str | None = None,
) -> Occurrence:
    occ = Occurrence(
        occurrence_id=occurrence_id(source, loc),
        source=source,
        format=fmt,
        location=dict(loc),
        expected=expected,
        defects=sorted(set(defects)),
        duplicate_of=duplicate_of,
    )
    p.occurrences.append(occ)
    return occ


def _when(cfg: ContactsConfig, rng: random.Random) -> datetime:
    span = int((cfg.end - cfg.start).total_seconds() // 60)
    base = datetime(cfg.start.year, cfg.start.month, cfg.start.day, tzinfo=UTC)
    return base + timedelta(minutes=rng.randint(0, span + 24 * 60 - 1))


# --- EML --------------------------------------------------------------------------


def _signature(v: _View, rng: random.Random) -> list[str]:
    lines = ["--", v.name()] if rng.random() < 0.5 else ["Best,", v.first or "", "", v.name()]
    if v.title and v.company:
        lines.append(rng.choice([f"{v.title} | {v.company}", f"{v.title}, {v.company}"]))
    elif v.title:
        lines.append(v.title)
    elif v.company:
        lines.append(v.company)
    if v.phone_text:
        lines.append(f"{rng.choice(['D:', 'm.', 'Tel', 'Phone:', 'Cell:'])} {v.phone_text}")
    if v.city and rng.random() < 0.4:
        lines.append(v.city)
    else:
        v.city = None  # not printed, so the source doesn't state it
    return lines


def _vcard(v: _View) -> str:
    lines = ["BEGIN:VCARD", "VERSION:3.0", f"N:{v.last};{v.first};;;", f"FN:{v.name()}"]
    if v.company:
        lines.append(f"ORG:{v.company}")
    if v.title:
        lines.append(f"TITLE:{v.title}")
    if v.phone:
        lines.append(f"TEL;TYPE=WORK,VOICE:{v.phone}")
    if v.email:
        lines.append(f"EMAIL;TYPE=INTERNET:{v.email}")
    if v.city:
        lines.append(f"ADR;TYPE=WORK:;;;{v.city};;;{v.country_text or ''}")
    lines.append("END:VCARD")
    return "\r\n".join(lines) + "\r\n"


def _build_eml(cfg: ContactsConfig, people: list[_Person], out: PackOutput, seed: int) -> None:
    rng = sub_rng(seed, PACK, "eml")
    senders = rng.sample(people, cfg.n_eml_senders)
    for i in range(1, cfg.n_eml + 1):
        p = rng.choice(senders)
        when = _when(cfg, rng)
        v = _view(
            p, rng, fields=("phone", "company", "title", "city"), p_drop=0.25, when=when.date()
        )
        subject = rng.choice(SUBJECTS).format(
            q=rng.randint(1, 4), thing=rng.choice(THINGS), day=when.strftime("%A")
        )
        name = f"{when:%Y%m%d}_{i:03d}.eml" if i % 3 else f"message_{i:04d}.eml"
        source = f"{PACK}/email/{name}"
        sender_email = p.old_email if "stale_company" in v.defects and p.old_email else p.email
        from_style = rng.choices(["normal", "bare", "last_first"], [7, 2, 1])[0]

        msg = EmailMessage()
        if from_style == "bare":
            msg["From"] = sender_email
        else:
            msg["From"] = f'"{v.name(from_style)}" <{sender_email}>'
        msg["To"] = "operations-desk@ourco.example"
        msg["Subject"] = subject
        msg["Date"] = format_datetime(when)
        msg["Message-ID"] = f"<{seed}.{i}@mail.generated.example>"

        body = [f"Hi {rng.choice(['team', 'there', 'all'])},", ""]
        body += rng.sample(BODY_LINES, rng.randint(1, 3))
        body += ["", *_signature(v, rng)]

        quoted: _Person | None = None
        if rng.random() < cfg.reply_chain_share:
            quoted = rng.choice([q for q in people if q is not p])
            qwhen = when - timedelta(days=rng.randint(1, 20), minutes=rng.randint(0, 600))
            qv = _view(
                quoted,
                rng,
                fields=("phone", "company", "title"),
                p_drop=0.4,
                p_typo=0,
                when=qwhen.date(),
            )
            body += ["", f"On {format_datetime(qwhen)}, {qv.name()} <{quoted.email}> wrote:"]
            qbody = [*rng.sample(BODY_LINES, 2), "", *_signature(qv, rng)]
            body += ["> " + line if line else ">" for line in qbody]
            q_exp = qv.expected()
            q_exp["email"] = quoted.email
            _occ(quoted, source, "eml", {"part": "quoted"}, q_exp, [*qv.defects, "quoted_reply"])
        text = "\n".join(body) + "\n"
        msg.set_content(text)
        if rng.random() < cfg.html_share:
            markup = (
                "<html><body>"
                + "".join(f"<p>{html.escape(line)}</p>" if line else "<br>" for line in body)
                + "</body></html>"
            )
            msg.add_alternative(markup, subtype="html")

        exp = v.expected()
        exp["email"] = sender_email
        defects = list(v.defects)
        if from_style == "last_first":
            defects.append("name_last_first")
        _occ(p, source, "eml", {"part": "body"}, exp, defects)

        if rng.random() < cfg.vcard_share:
            cv = _view(p, rng, canonical=True)
            fname = f"{_ascii(p.first)}_{_ascii(p.last)}.vcf"
            msg.add_attachment(
                _vcard(cv).encode("utf-8"), maintype="text", subtype="vcard", filename=fname
            )
            _occ(p, source, "eml", {"part": f"attachment:{fname}"}, cv.expected(), [])
        out.add(source, eml_bytes(msg, f"{seed}_{i}"))


# --- DOCX -----------------------------------------------------------------------------


def _build_docx(cfg: ContactsConfig, people: list[_Person], out: PackOutput, seed: int) -> None:
    rng = sub_rng(seed, PACK, "docx")
    kinds = ["contact_sheet", "attendees", "memo"]
    for i in range(1, cfg.n_docx + 1):
        kind = kinds[i % 3]
        group = rng.sample(people, rng.randint(3, 8))
        doc = Document()
        source = f"{PACK}/docx/{kind}_{i:02d}.docx"
        if kind == "contact_sheet":
            doc.add_heading(rng.choice(["Key Contacts", "Vendor & Partner Contact List",
                                        "Contacts: Building Ops"]), level=1)  # fmt: skip
            updated = _when(cfg, rng)
            doc.add_paragraph(f"Last updated {updated:%B %Y}. Internal use only.")
            table = doc.add_table(rows=1, cols=5)
            table.style = rng.choice(["Table Grid", "Light List Accent 1", "Light Shading"])
            for cell, h in zip(table.rows[0].cells, ["Name", "Title", "Company", "Phone",
                                                     "Email"], strict=True):  # fmt: skip
                cell.text = h
            for r, p in enumerate(group, 1):
                v = _view(
                    p,
                    rng,
                    fields=("email", "phone", "company", "title"),
                    p_drop=0.15,
                    when=updated.date().replace(day=1),
                )
                style = rng.choices(["normal", "last_first"], [3, 1])[0]
                row = table.add_row().cells
                values = [v.name(style), v.title or "", v.company or "", v.phone_text or "",
                          v.email or ""]  # fmt: skip
                for cell, value in zip(row, values, strict=True):
                    cell.text = value
                if rng.random() < 0.3:
                    row[0].paragraphs[0].runs[0].bold = True
                d = [*v.defects, "name_last_first"] if style == "last_first" else v.defects
                _occ(p, source, "docx", {"table": 0, "row": r}, v.expected(), d)
        elif kind == "attendees":
            thing = rng.choice(THINGS)
            doc.add_heading(f"Meeting: {thing[0].upper()}{thing[1:]}", level=2)
            met = _when(cfg, rng)
            doc.add_paragraph(f"Date: {met:%d %B %Y}")
            parts = []
            for p in group:
                v = _view(p, rng, fields=("company", "title"), p_drop=0.3, when=met.date())
                extra = ", ".join(x for x in (v.title, v.company) if x)
                parts.append(f"{v.name()} ({extra})" if extra else v.name())
                _occ(p, source, "docx", {"paragraph": 2, "mention": len(parts)}, v.expected(),
                     v.defects)  # fmt: skip
            doc.add_paragraph("Attendees: " + "; ".join(parts) + ".")
            doc.add_paragraph("Agenda: scope review, schedule, next steps.")
        else:
            doc.add_heading("Project Kickoff Memo", level=1)
            memo_date = _when(cfg, rng)
            doc.add_paragraph(f"Date: {memo_date:%B %d, %Y}")
            doc.add_paragraph(
                f"Re: {rng.choice(THINGS)}. The following people are the points of contact."
            )
            for p in group:
                v = _view(
                    p,
                    rng,
                    fields=("email", "phone", "company", "title"),
                    p_drop=0.2,
                    when=memo_date.date(),
                )
                bits = [v.name()]
                if v.title or v.company:
                    bits.append(", ".join(x for x in (v.title, v.company) if x))
                contact = " / ".join(x for x in (v.email, v.phone_text) if x)
                text = f"{rng.choice(['Primary contact', 'Contact', 'Escalation'])}: " + " — ".join(
                    bits
                )
                if contact:
                    text += f". Reach them at {contact}."
                para = doc.add_paragraph(text)
                if rng.random() < 0.5:
                    para.runs[0].italic = True
                _occ(p, source, "docx", {"paragraph": len(doc.paragraphs) - 1}, v.expected(),
                     v.defects)  # fmt: skip
        out.add(source, docx_bytes(doc))


# --- TXT / MD notes ----------------------------------------------------------------


def _build_notes(cfg: ContactsConfig, people: list[_Person], out: PackOutput, seed: int) -> None:
    rng = sub_rng(seed, PACK, "notes")
    for i in range(1, cfg.n_notes + 1):
        md = i % 2 == 0
        when = _when(cfg, rng)
        group = rng.sample(people, rng.randint(1, 3))
        lines: list[str] = []
        pending: list[tuple[_Person, int, dict[str, Any], list[str]]] = []
        if md:
            lines += [f"# {rng.choice(['Site visit', 'Call notes', 'Vendor sync'])}: "
                      f"{when:%Y-%m-%d}", "", "**Contacts**", ""]  # fmt: skip
        else:
            lines += [f"{when:%m/%d/%Y} call notes", ""]
        for p in group:
            v = _view(p, rng, fields=("email", "phone", "company", "title"), p_drop=0.35,
                      when=when.date())  # fmt: skip
            partial = md and rng.random() < 0.3
            name = v.name("partial") if partial else v.name()
            defects = [*v.defects, "partial_name"] if partial else v.defects
            if md:
                bits = [x for x in (v.title, v.company, v.phone_text, v.email) if x]
                lines.append(f"- {name}: " + " · ".join(bits) if bits else f"- {name}")
            else:
                sentence = f"Spoke with {name}"
                if v.company:
                    sentence += f" from {v.company}"
                if v.title:
                    sentence += f". {rng.choice(['Title:', 'Now'])} {v.title}"
                if v.phone_text:
                    sentence += f". Best number is {v.phone_text}"
                if v.email:
                    sentence += f", email {v.email}"
                lines += [sentence + ".", ""]
            pending.append((p, len(lines) if md else len(lines) - 1, v.expected(
                partial_name=partial), defects))  # fmt: skip
        lines += ["", "Follow-ups:" if not md else "## Follow-ups",
                  f"- {rng.choice(BODY_LINES)}"]  # fmt: skip
        ext = "md" if md else "txt"
        source = f"{PACK}/notes/note_{when:%Y%m%d}_{i:02d}.{ext}"
        for p, line_no, exp, defects in pending:
            _occ(p, source, ext, {"line": line_no}, exp, defects)  # type: ignore[arg-type]
        out.add(source, ("\n".join(lines) + "\n").encode("utf-8"))


# --- JSONL + log -------------------------------------------------------------------------


def _jsonl_record(v: _View, when: datetime, shape: int) -> dict[str, Any]:
    if shape == 0:
        return {"submitted_at": when.isoformat(), "name": v.name(), "email": v.email,
                "phone": v.phone_text, "org": v.company, "role": v.title}  # fmt: skip
    if shape == 1:
        return {"ts": int(when.timestamp()), "contact": {"firstName": v.first,
                "lastName": v.last, "emailAddress": v.email}, "company_name": v.company,
                "meta": {"phoneNumber": v.phone_text}}  # fmt: skip
    return {"data": {"full_name": v.name("upper_last_first"), "Email": v.email,
                     "tel": v.phone, "Company": v.company, "Country": v.country_text}}  # fmt: skip


def _build_jsonl_and_log(
    cfg: ContactsConfig, people: list[_Person], out: PackOutput, seed: int
) -> list[list[str]]:
    rng = sub_rng(seed, PACK, "jsonl")
    source = f"{PACK}/exports/webform_submissions.jsonl"
    lines: list[str] = []
    noise_at = set(rng.sample(range(cfg.n_jsonl_lines), cfg.n_jsonl_noise))
    first_seen: dict[str, str] = {}
    groups: dict[str, list[str]] = {}
    dup_budget = cfg.n_jsonl_dups
    for n in range(cfg.n_jsonl_lines):
        when = _when(cfg, rng)
        if n in noise_at:
            lines.append(rng.choice(['{"event": "heartbeat", "ts": ', '{"test": true}',
                                     '{"event":"form_view","path":"/contact"}']))  # fmt: skip
            continue
        repeat = dup_budget > 0 and first_seen and rng.random() < 0.12
        if repeat:
            rid = rng.choice(sorted(first_seen))
            p = next(x for x in people if x.record_id == rid)
            dup_budget -= 1
        else:
            p = rng.choice(people)
        v = _view(
            p, rng, fields=("email", "phone", "company", "country"), p_drop=0.25, when=when.date()
        )
        shape = rng.randrange(3)
        lines.append(json.dumps(_jsonl_record(v, when, shape), ensure_ascii=bool(shape == 1)))
        exp = v.expected()
        if shape != 2:  # only the third shape carries a country
            exp.pop("country", None)
        defects = list(v.defects)
        if shape == 2:
            defects.append("name_upper_last_first")
        loc = {"line": len(lines)}
        oid = occurrence_id(source, loc)
        if repeat and p.record_id in first_seen:
            defects += ["duplicate", "repeat_submission"]
            _occ(p, source, "jsonl", loc, exp, defects, duplicate_of=first_seen[p.record_id])
            groups.setdefault(p.record_id, [first_seen[p.record_id]]).append(oid)
        else:
            _occ(p, source, "jsonl", loc, exp, defects)
            first_seen.setdefault(p.record_id, oid)
    out.add(source, ("\n".join(lines) + "\n").encode("utf-8"))

    # Application log: three logger formats interleaved with noise.
    rng = sub_rng(seed, PACK, "log")
    source = f"{PACK}/exports/signups.log"
    log: list[str] = []
    when = datetime(2024, 1, 2, 8, 0, tzinfo=UTC)
    for _ in range(cfg.n_log_lines):
        when += timedelta(minutes=rng.randint(3, 4000))
        if rng.random() >= cfg.log_contact_share:
            log.append(rng.choice([
                f"[{when:%Y-%m-%d %H:%M:%S}] DEBUG cache hit key=lead:{rng.randint(1000, 9999)}",
                f"{when:%Y-%m-%dT%H:%M:%SZ} INFO  [scheduler] job=sync_crm status=ok "
                f"duration_ms={rng.randint(80, 4000)}",
                f"{when:%Y-%m-%d %H:%M:%S},{rng.randint(100, 999)} ERROR smtp: timeout "
                "connecting to relay",
            ]))  # fmt: skip
            continue
        p = rng.choice(people)
        v = _view(p, rng, fields=("email", "phone", "company"), p_drop=0.3, when=when.date())
        fmt = rng.randrange(3)
        if fmt == 0:
            kv = [f'name="{v.name()}"']
            if v.email:
                kv.append(f"email={v.email}")
            if v.company:
                kv.append(f'company="{v.company}"')
            if v.phone_text:
                kv.append(f'phone="{v.phone_text}"')
            log.append(f"{when:%Y-%m-%dT%H:%M:%SZ} INFO  [webform] new lead " + " ".join(kv))
        elif fmt == 1:
            payload = {"first": v.first, "last": v.last, "mail": v.email, "org": v.company}
            payload = {k: val for k, val in payload.items() if val}
            v.phone = None
            log.append(f"{when:%Y-%m-%d %H:%M:%S},{rng.randint(100, 999)} INFO lead_capture: "
                       f"payload={json.dumps(payload, ensure_ascii=False)}")  # fmt: skip
        else:
            v.company = None
            v.phone = None
            log.append(f"[{when:%Y-%m-%d %H:%M:%S}] WARN crm: contact {v.name()} <{v.email or ''}>"
                       " missing owner, assigned to queue")  # fmt: skip
        _occ(p, source, "log", {"line": len(log)}, v.expected(), v.defects)
    out.add(source, ("\n".join(log) + "\n").encode("utf-8"))
    return list(groups.values())


# --- CRM CSV ---------------------------------------------------------------------------


def _build_crm(
    cfg: ContactsConfig, people: list[_Person], out: PackOutput, seed: int
) -> list[list[str]]:
    rng = sub_rng(seed, PACK, "crm")
    members = [p for p in people if rng.random() < cfg.crm_share or not p.occurrences]
    header = ["Contact ID", "First Name", "Last Name", "Email", "Phone", "Account Name",
              "Title", "Mailing City", "Mailing Country"]  # fmt: skip
    entries: list[tuple[_Person, list[str], dict[str, Any], list[str]]] = []
    for p in members:
        v = _view(p, rng, p_drop=0.1, when=cfg.end)  # the export is current
        row = [f"CRM-{int(p.record_id[4:]) * 7 + 1000:06d}", v.first or "", v.last or "",
               v.email or "", v.phone_text or "", v.company or "", v.title or "", v.city or "",
               v.country_text or ""]  # fmt: skip
        entries.append((p, row, v.expected(), v.defects))
    rng.shuffle(entries)
    for j, i in enumerate(rng.sample(range(len(entries)), cfg.n_crm_dups)):
        p, row, exp, defects = entries[i]
        if j % 2 == 0:
            entries.insert(rng.randint(i + 1, len(entries)),
                           (p, list(row), dict(exp), [*defects, "duplicate",
                                                      "exact_duplicate_row"]))  # fmt: skip
        else:
            v = _view(p, rng, p_nick=0.9, p_drop=0.3, when=cfg.end)
            row2 = [f"CRM-{int(p.record_id[4:]) * 7 + 5000:06d}", v.first or "", v.last or "",
                    v.email or "", v.phone_text or "", v.company or "", v.title or "",
                    v.city or "", v.country_text or ""]  # fmt: skip
            entries.insert(rng.randint(i + 1, len(entries)),
                           (p, row2, v.expected(), [*v.defects, "duplicate",
                                                    "near_duplicate_row"]))  # fmt: skip
    source = f"{PACK}/crm_export.csv"
    rows = [header]
    first_seen: dict[str, str] = {}
    groups: dict[str, list[str]] = {}
    for p, row, exp, defects in entries:
        rows.append(row)
        loc = {"row": len(rows)}
        oid = occurrence_id(source, loc)
        if "duplicate" in defects and p.record_id in first_seen:
            _occ(p, source, "csv", loc, exp, defects, duplicate_of=first_seen[p.record_id])
            groups.setdefault(p.record_id, [first_seen[p.record_id]]).append(oid)
        else:
            _occ(p, source, "csv", loc, exp, [d for d in defects if "duplicate" not in d])
            first_seen.setdefault(p.record_id, oid)
    out.add(source, csv_bytes(rows, encoding="utf-8"))
    return list(groups.values())


# --- entry point ----------------------------------------------------------------------


def generate(seed: int, cfg: ContactsConfig | None = None) -> PackOutput:
    cfg = cfg or ContactsConfig()
    out = PackOutput()
    people = _make_people(cfg, sub_rng(seed, PACK, "people"))
    _build_eml(cfg, people, out, seed)
    _build_docx(cfg, people, out, seed)
    _build_notes(cfg, people, out, seed)
    dup_groups = _build_jsonl_and_log(cfg, people, out, seed)
    dup_groups += _build_crm(cfg, people, out, seed)

    aliases: dict[str, set[str]] = {c.company_id: set() for c in COMPANIES}
    by_name = {c.name: c for c in COMPANIES}
    for p in people:
        for o in p.occurrences:
            shown = o.expected.get("company")
            # Attribute the alias to the company actually meant (old employer if stale).
            meant = p.old_company if "stale_company" in o.defects and p.old_company else p.company
            if shown and shown not in by_name:
                aliases[meant.company_id].add(shown)
    entities = [
        Entity(
            entity_id=c.company_id,
            kind="company",
            canonical_name=c.name,
            aliases=sorted(aliases[c.company_id]),
            lookalike_of=c.lookalike_of,
        )
        for c in COMPANIES
    ]
    out.ground_truth = GroundTruth(
        pack=PACK,
        seed=seed,
        generator_version=GENERATOR_VERSION,
        entities=entities,
        records=[
            TrueRecord(
                record_id=p.record_id,
                entity_id=p.company.company_id,
                truth=p.truth().model_dump(mode="json"),
                occurrences=p.occurrences,
            )
            for p in people
        ],
        duplicate_groups=dup_groups,
    )
    return out
