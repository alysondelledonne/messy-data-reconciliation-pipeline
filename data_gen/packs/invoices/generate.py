"""Invoices pack: 300 true invoices spread across PDFs, two CSV exports and an XLSX report.

Sources and the mess planted in each:
- `pdf/`: about 150 vendor PDFs in seven layouts. About 18 are scanned (image only).
  Planted: duplicates (exact resends, stamped re-issues, re-scans) and totals that
  don't add up.
- `ap_register.csv` (legacy AP): Latin-1, a report preamble, mixed date formats
  including dates with no year, `$1,234.50` amounts, uppercase vendor names cut off
  at 30 characters, reformatted invoice numbers, trailing blank rows.
- `procurement_export.csv`: UTF-8 with a BOM, a different schema, vendor alias
  spellings, missing invoice numbers, duplicate rows.
- `ap_aging.xlsx`: open (unpaid) items. Title rows above the header, supplier shown
  only on the first row of each group, subtotal rows, dates stored as serial
  numbers, stray notes. A few rows are stale: the invoice was paid but is still
  listed as open.
"""

import random
from dataclasses import dataclass
from datetime import date, timedelta
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

from data_gen import GENERATOR_VERSION
from data_gen.corpus import PackOutput, sub_rng
from data_gen.ground_truth import Entity, GroundTruth, Occurrence, TrueRecord, occurrence_id
from data_gen.mess import (
    excel_serial,
    format_amount,
    format_date,
    money,
    org_variant,
    random_org_variant,
    transpose_digits,
)
from data_gen.packs.invoices.pdf_layouts import L, PrintedInvoice, PrintedLine, render_invoice
from data_gen.packs.invoices.vocab import CATALOG, VENDORS, Vendor
from data_gen.writers.files import csv_bytes, xlsx_bytes
from data_gen.writers.scan import scan_pdf
from pipeline.domains.invoices.schema import Invoice, LineItem

PACK = "invoices"


@dataclass(frozen=True)
class InvoicesConfig:
    n_records: int = 300
    n_pdf: int = 150
    n_scanned: int = 18
    n_resend: int = 2
    n_reissue: int = 2
    n_rescan: int = 2
    n_total_mismatch: int = 8
    csv_a_share: float = 0.85
    csv_b_share: float = 0.70
    csv_b_dup_rows: int = 8
    csv_b_missing_number_share: float = 0.15
    stale_aging_rows: int = 5
    long_invoice_share: float = 0.12
    start: date = date(2023, 1, 1)
    end: date = date(2025, 12, 31)


@dataclass
class _Rec:
    record_id: str
    vendor: Vendor
    invoice: Invoice
    occurrences: list[Occurrence]


def _number(vendor: Vendor, n: int, d: date) -> str:
    return vendor.number_pattern.format(n=n, yy=f"{d.year % 100:02d}", yyyy=d.year)


def _lines(vendor: Vendor, rng: random.Random, long: bool) -> list[LineItem]:
    catalog = CATALOG[vendor.trade]
    count = rng.randint(20, 45) if long else rng.choices(range(1, 9), [5, 6, 5, 4, 3, 2, 1, 1])[0]
    lines = []
    for _ in range(count):
        desc, unit, lo, hi = rng.choice(catalog)
        if unit == "hr":
            qty = Decimal(rng.randint(2, 32)) / 2
        elif vendor.trade == "supplies":
            qty = Decimal(rng.randint(1, 24))
        else:
            qty = Decimal(rng.choices([1, 2, 3, 4], [8, 3, 1, 1])[0])
        price = money(rng.uniform(lo, hi))
        lines.append(
            LineItem(description=desc, quantity=qty, unit_price=price, amount=money(qty * price))
        )
    return lines


def _make_records(cfg: InvoicesConfig, rng: random.Random) -> list[_Rec]:
    span = (cfg.end - cfg.start).days
    dates = sorted(cfg.start + timedelta(days=rng.randint(0, span)) for _ in range(cfg.n_records))
    counters = {v.vendor_id: rng.randint(100, 4000) for v in VENDORS}
    weights = [v.weight for v in VENDORS]
    recs: list[_Rec] = []
    for i, d in enumerate(dates, 1):
        vendor = rng.choices(VENDORS, weights)[0]
        counters[vendor.vendor_id] += rng.randint(1, 9)
        lines = _lines(vendor, rng, rng.random() < cfg.long_invoice_share)
        subtotal = money(sum(li.amount for li in lines))
        tax = money(subtotal * vendor.tax_rate)
        recent = (cfg.end - d).days < 120
        unpaid = rng.random() < (0.75 if recent else 0.12)
        po = f"PO-{d.year % 100:02d}{rng.randint(10000, 99999)}" if rng.random() < 0.6 else None
        inv = Invoice(
            invoice_number=_number(vendor, counters[vendor.vendor_id], d),
            vendor_name=vendor.name,
            invoice_date=d,
            due_date=d + timedelta(days=vendor.terms_days),
            currency=vendor.currency,
            po_number=po,
            line_items=lines,
            subtotal=subtotal,
            tax=tax,
            total=subtotal + tax,
            payment_status="unpaid" if unpaid else "paid",
        )
        recs.append(_Rec(f"inv-{i:04d}", vendor, inv, []))
    return recs


# --- PDFs ------------------------------------------------------------------------


def _slug(s: str) -> str:
    return "".join(ch if ch.isalnum() else "-" for ch in s).strip("-")


def _printed(rec: _Rec) -> PrintedInvoice:
    inv, v = rec.invoice, rec.vendor
    assert inv.invoice_number and inv.due_date and inv.subtotal is not None
    assert inv.tax is not None
    return PrintedInvoice(
        vendor_name=v.name,
        vendor_address=v.address,
        invoice_number=inv.invoice_number,
        invoice_date=inv.invoice_date,
        due_date=inv.due_date,
        terms_days=v.terms_days,
        po_number=inv.po_number,
        currency=inv.currency,
        tax_rate=v.tax_rate,
        lines=[
            PrintedLine(li.description, li.quantity, li.unit_price or Decimal(0), li.amount)
            for li in inv.line_items
        ],
        subtotal=inv.subtotal,
        tax=inv.tax,
        total=inv.total,
    )


def _plant_mismatch(p: PrintedInvoice, rng: random.Random, has_subtotal: bool) -> str:
    """Corrupt the printed totals. Returns the defect name."""
    if has_subtotal and len(p.lines) > 1 and rng.random() < 0.5:
        # A line was left out of the subtotal; tax and total follow the wrong subtotal.
        p.subtotal = money(p.subtotal - p.lines[-1].amount)
        p.tax = money(p.subtotal * p.tax_rate)
        p.total = p.subtotal + p.tax
        return "subtotal_mismatch"
    if rng.random() < 0.5:
        p.total = money(p.total + Decimal(rng.choice([10, 25, 50, 100, -20])))
    else:
        bumped = transpose_digits(f"{p.total:.2f}", rng)
        p.total = Decimal(bumped) if bumped != f"{p.total:.2f}" else money(p.total + 1)
    return "total_mismatch"


def _pdf_expected(p: PrintedInvoice, layout: str) -> dict[str, Any]:
    spec = L[layout]
    lines = []
    for li in p.lines:
        d: dict[str, Any] = {
            "description": li.description,
            "quantity": str(li.quantity),
            "amount": f"{li.amount:.2f}",
        }
        if spec.show_unit_price:
            d["unit_price"] = f"{li.unit_price:.2f}"
        lines.append(d)
    exp: dict[str, Any] = {
        "vendor_name": p.vendor_name,
        "invoice_number": p.invoice_number,
        "invoice_date": p.invoice_date.isoformat(),
        "line_items": lines,
        "tax": f"{p.tax:.2f}",
        "total": f"{p.total:.2f}",
    }
    # Only claim a currency the page actually shows: a symbol, a code, or an "amounts in" note.
    if spec.money_style in ("symbol", "code", "eu", "eu_code") or spec.currency_note:
        exp["currency"] = p.currency
    if spec.print_due_date:
        exp["due_date"] = p.due_date.isoformat()
    if spec.print_subtotal:
        exp["subtotal"] = f"{p.subtotal:.2f}"
    if p.po_number:
        exp["po_number"] = p.po_number
    return exp


def _pdf_name(rec: _Rec, k: int, rng: random.Random) -> str:
    num = _slug(rec.invoice.invoice_number or "")
    style = rng.randrange(4)
    if style == 0:
        return f"{_slug(rec.vendor.name.split()[0]).lower()}_{num}.pdf"
    if style == 1:
        return f"Invoice_{num}.pdf"
    if style == 2:
        return f"inv-{rec.invoice.invoice_date:%Y%m%d}-{k:03d}.pdf"
    return f"{num} {rec.vendor.name.split()[0]}.pdf"


def _build_pdfs(
    cfg: InvoicesConfig, recs: list[_Rec], out: PackOutput, seed: int
) -> list[list[str]]:
    rng = sub_rng(seed, PACK, "pdf")
    chosen = sorted(rng.sample(range(len(recs)), cfg.n_pdf))
    scanned = set(rng.sample(chosen, cfg.n_scanned))
    mismatch = set(rng.sample(chosen, cfg.n_total_mismatch))
    dup_groups: list[list[str]] = []
    text_occ: list[tuple[_Rec, Occurrence, bytes, PrintedInvoice]] = []
    used_names: set[str] = set()

    def unique(name: str) -> str:
        stem, ext = name.rsplit(".", 1)
        candidate, n = name, 2
        while candidate in used_names:
            candidate, n = f"{stem}_{n}.{ext}", n + 1
        used_names.add(candidate)
        return candidate

    for k, idx in enumerate(chosen, 1):
        rec = recs[idx]
        layout = rec.vendor.layout
        p = _printed(rec)
        defects: list[str] = []
        if idx in mismatch:
            defects.append(_plant_mismatch(p, rng, L[layout].print_subtotal))
        pdf, pages = render_invoice(p, L[layout])
        if pages > 1:
            defects.append("multi_page")
        fmt: Any = "pdf"
        if idx in scanned:
            pdf, _ = scan_pdf(pdf, rng)
            fmt = "pdf_scanned"
            scan_day = rec.invoice.invoice_date + timedelta(days=rng.randint(1, 12))
            name = unique(f"scan_{scan_day:%Y%m%d}_{k:03d}.pdf")
        else:
            name = unique(_pdf_name(rec, k, rng))
        source = f"{PACK}/pdf/{name}"
        loc = {"pages": pages}
        occ = Occurrence(
            occurrence_id=occurrence_id(source, loc),
            source=source,
            format=fmt,
            location=loc,
            layout=layout,
            expected=_pdf_expected(p, layout),
            defects=defects,
        )
        rec.occurrences.append(occ)
        out.add(source, pdf)
        if fmt == "pdf":
            text_occ.append((rec, occ, pdf, p))

    # Planted duplicates, each copying a distinct text PDF.
    picks = rng.sample(text_occ, cfg.n_resend + cfg.n_reissue + cfg.n_rescan)
    for j, (rec, orig, pdf, p) in enumerate(picks):
        stem = orig.source.rsplit("/", 1)[1].rsplit(".", 1)[0]
        fmt = "pdf"
        if j < cfg.n_resend:
            kind, name, data = "exact_copy", unique(f"{stem} (1).pdf"), pdf
        elif j < cfg.n_resend + cfg.n_reissue:
            p.stamp = "COPY"
            data, _ = render_invoice(p, L[orig.layout or "classic"])
            kind, name = "reissued_copy", unique(f"{stem}_reissued.pdf")
        else:
            data, _ = scan_pdf(pdf, rng)
            kind, fmt = "rescan", "pdf_scanned"
            name = unique(f"scan_{rec.invoice.invoice_date:%Y%m%d}_dup{j:02d}.pdf")
        source = f"{PACK}/pdf/{name}"
        dup = Occurrence(
            occurrence_id=occurrence_id(source, orig.location),
            source=source,
            format=fmt,
            location=orig.location,
            layout=orig.layout,
            expected=orig.expected,
            defects=[*orig.defects, "duplicate", kind],
            duplicate_of=orig.occurrence_id,
        )
        rec.occurrences.append(dup)
        out.add(source, data)
        dup_groups.append([orig.occurrence_id, dup.occurrence_id])
    return dup_groups


# --- CSV A: legacy AP register ----------------------------------------------------------

A_DATE_STYLES = ["us_short", "iso", "month_day", "mon_d_yyyy", "us_long"]
A_DATE_WEIGHTS = [4, 3, 1, 2, 2]


def _legacy_master_name(v: Vendor, rng: random.Random) -> str:
    """How the old AP system's vendor master spells this vendor."""
    roll = rng.random()
    if roll < 0.6:
        return v.name.upper()[:30].rstrip()
    if roll < 0.85:
        return v.name
    return random_org_variant(v.name, rng)


def _build_csv_a(cfg: InvoicesConfig, recs: list[_Rec], out: PackOutput, seed: int) -> set[str]:
    rng = sub_rng(seed, PACK, "csv_a")
    master = {v.vendor_id: _legacy_master_name(v, rng) for v in VENDORS}
    members = [r for r in recs if rng.random() < cfg.csv_a_share]
    # Every record must appear somewhere; records that are in no other source land here.
    pdf_ids = {r.record_id for r in recs if r.occurrences}
    b_rng = sub_rng(seed, PACK, "csv_b_membership")
    in_b = {r.record_id for r in recs if b_rng.random() < cfg.csv_b_share}
    member_ids = {r.record_id for r in members}
    for r in recs:
        if r.record_id not in member_ids | pdf_ids | in_b:
            members.append(r)
    members.sort(key=lambda r: (r.invoice.invoice_date, rng.random()))

    header = ["Vendor Name", "Invoice #", "Inv Date", "Period", "Amount", "Curr", "Status",
              "Paid Date", "Memo"]  # fmt: skip
    preamble = ["AP INVOICE REGISTER", "Exported 01/05/2026 06:00 by APSYS v4.2"]
    rows = [header]
    source = f"{PACK}/ap_register.csv"
    for r in members:
        inv, v = r.invoice, r.vendor
        assert inv.invoice_number
        defects: list[str] = []
        name = master[v.vendor_id]
        if rng.random() < 0.08:
            name = random_org_variant(v.name, rng)
        number = inv.invoice_number
        if rng.random() < 0.1:
            stripped = "".join(ch for ch in number if ch.isalnum())
            if stripped != number:
                number, defects = stripped, [*defects, "number_reformatted"]
        style = rng.choices(A_DATE_STYLES, A_DATE_WEIGHTS)[0]
        if style == "month_day":
            defects.append("date_without_year")
        if inv.currency == "USD":
            amt_style = rng.choices(["symbol", "grouped", "plain"], [5, 3, 2])[0]
            curr = "USD" if rng.random() < 0.5 else ""
        else:
            amt_style = rng.choice(["code", "plain"])
            curr = inv.currency if amt_style == "plain" or rng.random() < 0.5 else ""
        paid = inv.payment_status == "paid"
        status = rng.choice(["PAID", "Paid", "P"]) if paid else rng.choice(["OPEN", "Open", "O"])
        paid_date = (
            format_date(inv.invoice_date + timedelta(days=rng.randint(10, 60)), style)
            if paid
            else ""
        )
        memo = f"PO {inv.po_number}" if inv.po_number and rng.random() < 0.5 else ""
        rows.append([
            name, number, format_date(inv.invoice_date, style), f"{inv.invoice_date:%Y-%m}",
            format_amount(inv.total, amt_style, inv.currency), curr, status, paid_date, memo,
        ])  # fmt: skip
        expected: dict[str, Any] = {
            "vendor_name": name,
            "invoice_number": number,
            "invoice_date": inv.invoice_date.isoformat(),
            "total": f"{inv.total:.2f}",
            "payment_status": inv.payment_status,
        }
        if curr or amt_style in ("symbol", "code"):
            expected["currency"] = inv.currency
        if memo:
            expected["po_number"] = inv.po_number
        if name != v.name:
            defects.append("vendor_alias")
        loc = {"row": len(preamble) + len(rows)}
        r.occurrences.append(Occurrence(
            occurrence_id=occurrence_id(source, loc), source=source, format="csv",
            location=loc, expected=expected, defects=defects,
        ))  # fmt: skip
    out.add(source, csv_bytes(rows, encoding="latin-1", preamble=preamble, trailing_blank_rows=3))
    return in_b


# --- CSV B: procurement export --------------------------------------------------------


def _build_csv_b(
    cfg: InvoicesConfig, recs: list[_Rec], in_b: set[str], out: PackOutput, seed: int
) -> list[list[str]]:
    rng = sub_rng(seed, PACK, "csv_b")
    master = {v.vendor_id: random_org_variant(v.name, rng) for v in VENDORS}
    header = ["po_id", "supplier", "supplier_invoice_ref", "invoice_dt", "net_amount",
              "tax_amount", "gross_amount", "currency_code", "approval_state",
              "line_count"]  # fmt: skip
    entries: list[tuple[_Rec, list[str], dict[str, Any], list[str]]] = []
    for r in recs:
        if r.record_id not in in_b:
            continue
        inv, v = r.invoice, r.vendor
        assert inv.subtotal is not None and inv.tax is not None
        name = master[v.vendor_id]
        if rng.random() < 0.2:
            name = random_org_variant(v.name, rng)
        defects = ["vendor_alias"] if name != v.name else []
        ref = inv.invoice_number or ""
        if rng.random() < cfg.csv_b_missing_number_share:
            ref, defects = "", [*defects, "missing_invoice_number"]
        dt = format_date(inv.invoice_date, rng.choice(["iso", "iso_time"]))
        state = (
            "CLOSED"
            if inv.payment_status == "paid"
            else rng.choice(["APPROVED", "PENDING_RECEIPT"])
        )
        row = [inv.po_number or "", name, ref, dt, f"{inv.subtotal:.2f}", f"{inv.tax:.2f}",
               f"{inv.total:.2f}", inv.currency, state, str(len(inv.line_items))]  # fmt: skip
        exp: dict[str, Any] = {
            "vendor_name": name,
            "invoice_date": inv.invoice_date.isoformat(),
            "subtotal": f"{inv.subtotal:.2f}",
            "tax": f"{inv.tax:.2f}",
            "total": f"{inv.total:.2f}",
            "currency": inv.currency,
        }
        if ref:
            exp["invoice_number"] = ref
        if inv.po_number:
            exp["po_number"] = inv.po_number
        entries.append((r, row, exp, defects))
    rng.shuffle(entries)

    # Duplicate rows: half exact, half near-duplicates; inserted at random later positions.
    for j, i in enumerate(rng.sample(range(len(entries)), cfg.csv_b_dup_rows)):
        r, row, exp, defects = entries[i]
        row2, exp2 = list(row), dict(exp)
        kind = "exact_duplicate_row"
        if j % 2:
            kind = "near_duplicate_row"
            row2[1] = org_variant(row[1], rng.choice(["upper", "spacing"]), rng)
            row2[3] = format_date(
                r.invoice.invoice_date, "iso_time" if "T" not in row[3] else "iso"
            )
            exp2["vendor_name"] = row2[1]
        pos = rng.randint(i + 1, len(entries))
        entries.insert(pos, (r, row2, exp2, [*defects, "duplicate", kind]))

    source = f"{PACK}/procurement_export.csv"
    rows = [header]
    first_seen: dict[str, str] = {}
    groups: dict[str, list[str]] = {}
    for r, row, exp, defects in entries:
        rows.append(row)
        loc = {"row": len(rows)}
        oid = occurrence_id(source, loc)
        is_dup = "duplicate" in defects
        r.occurrences.append(Occurrence(
            occurrence_id=oid, source=source, format="csv", location=loc, expected=exp,
            defects=defects, duplicate_of=first_seen.get(r.record_id) if is_dup else None,
        ))  # fmt: skip
        if is_dup:
            groups.setdefault(r.record_id, [first_seen[r.record_id]]).append(oid)
        else:
            first_seen[r.record_id] = oid
    out.add(source, csv_bytes(rows, encoding="utf-8", bom=True, line_terminator="\n"))
    return list(groups.values())


# --- XLSX: AP aging report --------------------------------------------------------------


def _build_aging(cfg: InvoicesConfig, recs: list[_Rec], out: PackOutput, seed: int) -> None:
    rng = sub_rng(seed, PACK, "aging")
    a_rng = sub_rng(seed, PACK, "csv_a")
    master = {v.vendor_id: _legacy_master_name(v, a_rng) for v in VENDORS}  # same AP system
    as_of = cfg.end
    open_items = [r for r in recs if r.invoice.payment_status == "unpaid"]
    paid = [r for r in recs if r.invoice.payment_status == "paid"]
    stale = {r.record_id for r in rng.sample(paid, cfg.stale_aging_rows)}
    items = open_items + [r for r in paid if r.record_id in stale]
    items.sort(key=lambda r: (master[r.vendor.vendor_id].lower(), r.invoice.invoice_date))

    wb = Workbook()
    notes = wb.active
    assert notes is not None
    notes.title = "Notes"
    notes["A1"] = "Aging run for month-end close."
    notes["A3"] = "Credit memos not netted. Ask AP about the Lakeshore items."
    ws = wb.create_sheet("Open Items")
    ws.merge_cells("A1:H1")
    ws["A1"] = "Accounts Payable Aging Report"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"As of {as_of:%d-%b-%Y}"
    ws["A3"] = "Entity: (all)    Currency: mixed (not converted)"
    header = ["Supplier", "Document", "Doc Date", "Due Date", "Amount", "Cur", "Days Past Due",
              "Bucket"]  # fmt: skip
    ws.append([])
    ws.append(header)
    for cell in ws[5]:
        cell.font = Font(bold=True)
    ws["J7"] = "check w/ AP re: dup?"

    source = f"{PACK}/ap_aging.xlsx"
    group_name: str | None = None
    group_total = Decimal(0)
    grand = Decimal(0)

    def close_group() -> None:
        nonlocal group_total
        if group_name is not None:
            ws.append([f"Total {group_name}", None, None, None, float(group_total)])
            ws.cell(ws.max_row, 1).font = Font(bold=True)
        group_total = Decimal(0)

    for r in items:
        inv, v = r.invoice, r.vendor
        assert inv.due_date and inv.invoice_number
        name = master[v.vendor_id]
        first_in_group = name != group_name
        if first_in_group:
            close_group()
            group_name = name
        overdue = max(0, (as_of - inv.due_date).days)
        bucket = ("Current" if overdue == 0 else "1-30" if overdue <= 30 else "31-60"
                  if overdue <= 60 else "61-90" if overdue <= 90 else "90+")  # fmt: skip
        roll = rng.random()
        doc_date: Any
        defects: list[str] = []
        if roll < 0.7:
            doc_date = inv.invoice_date
        elif roll < 0.9:
            doc_date, defects = excel_serial(inv.invoice_date), ["date_as_serial"]
        else:
            doc_date = format_date(inv.invoice_date, "slash_ymd")
        ws.append([name if first_in_group else None, inv.invoice_number, doc_date, inv.due_date,
                   float(inv.total), inv.currency, overdue, bucket])  # fmt: skip
        row = ws.max_row
        if isinstance(doc_date, date):
            ws.cell(row, 3).number_format = "dd-mmm-yy"
        ws.cell(row, 4).number_format = "yyyy-mm-dd"
        group_total += inv.total
        grand += inv.total
        if not first_in_group:
            defects.append("supplier_fill_down")
        if name != v.name:
            defects.append("vendor_alias")
        if inv.payment_status == "paid":
            defects.append("stale_status")
        loc: dict[str, int | str] = {"sheet": "Open Items", "row": row}
        r.occurrences.append(Occurrence(
            occurrence_id=occurrence_id(source, loc), source=source, format="xlsx",
            location=loc, defects=defects,
            expected={
                "vendor_name": name,
                "invoice_number": inv.invoice_number,
                "invoice_date": inv.invoice_date.isoformat(),
                "due_date": inv.due_date.isoformat(),
                "total": f"{inv.total:.2f}",
                "currency": inv.currency,
                "payment_status": "unpaid",
            },
        ))  # fmt: skip
    close_group()
    ws.append([])
    ws.append(["GRAND TOTAL (mixed currency)", None, None, None, float(grand)])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    for col, w in zip("ABCDEFGH", [34, 16, 12, 12, 12, 6, 14, 10], strict=True):
        ws.column_dimensions[col].width = w
    out.add(source, xlsx_bytes(wb))


# --- entry point ---------------------------------------------------------------------


def generate(seed: int, cfg: InvoicesConfig | None = None) -> PackOutput:
    cfg = cfg or InvoicesConfig()
    out = PackOutput()
    recs = _make_records(cfg, sub_rng(seed, PACK, "records"))
    dup_groups = _build_pdfs(cfg, recs, out, seed)
    in_b = _build_csv_a(cfg, recs, out, seed)
    dup_groups += _build_csv_b(cfg, recs, in_b, out, seed)
    _build_aging(cfg, recs, out, seed)

    aliases: dict[str, set[str]] = {v.vendor_id: set() for v in VENDORS}
    for r in recs:
        for o in r.occurrences:
            if o.expected["vendor_name"] != r.vendor.name:
                aliases[r.vendor.vendor_id].add(o.expected["vendor_name"])
    entities = [
        Entity(
            entity_id=v.vendor_id,
            kind="vendor",
            canonical_name=v.name,
            aliases=sorted(aliases[v.vendor_id]),
            lookalike_of=v.lookalike_of,
        )
        for v in VENDORS
    ]
    out.ground_truth = GroundTruth(
        pack=PACK,
        seed=seed,
        generator_version=GENERATOR_VERSION,
        entities=entities,
        records=[
            TrueRecord(
                record_id=r.record_id,
                entity_id=r.vendor.vendor_id,
                truth=r.invoice.model_dump(mode="json"),
                occurrences=r.occurrences,
            )
            for r in recs
        ],
        duplicate_groups=dup_groups,
    )
    return out
