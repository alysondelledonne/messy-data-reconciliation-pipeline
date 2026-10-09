"""The invoices corpus holds the mess it claims to, and its ground truth matches the files."""

import csv
import io
from decimal import Decimal
from typing import Any

import pdfplumber
import pytest
from openpyxl import load_workbook

from data_gen.ground_truth import GroundTruth, Occurrence
from data_gen.packs.invoices.generate import InvoicesConfig
from pipeline.domains.invoices.schema import Invoice

CFG = InvoicesConfig()


def _occs(gt: GroundTruth, fmt: str | None = None) -> list[tuple[dict[str, Any], Occurrence]]:
    return [
        (r.truth, o) for r in gt.records for o in r.occurrences if fmt is None or o.format == fmt
    ]


def test_truth_and_expected_values_fit_the_schema(invoices_gt: GroundTruth) -> None:
    for r in invoices_gt.records:
        inv = Invoice.model_validate(r.truth)
        assert inv.subtotal == sum((li.amount for li in inv.line_items), Decimal(0))
        assert inv.tax is not None and inv.total == inv.subtotal + inv.tax
        for o in r.occurrences:
            Invoice.model_validate({**r.truth, **o.expected})  # partial values are well-typed


def test_pdf_counts(invoices_gt: GroundTruth) -> None:
    pdfs = [o for _, o in _occs(invoices_gt) if o.format in ("pdf", "pdf_scanned")]
    dups = CFG.n_resend + CFG.n_reissue + CFG.n_rescan
    assert len(pdfs) == CFG.n_pdf + dups
    assert sum(o.format == "pdf_scanned" for o in pdfs) == CFG.n_scanned + CFG.n_rescan
    assert sum("duplicate" in o.defects for o in pdfs) == dups
    originals = [o for o in pdfs if "duplicate" not in o.defects]
    mismatched = [o for o in originals if {"total_mismatch", "subtotal_mismatch"} & set(o.defects)]
    assert len(mismatched) == CFG.n_total_mismatch
    assert any("multi_page" in o.defects for o in pdfs)
    assert len({o.layout for o in pdfs}) >= 6


def test_planted_mismatches_do_not_add_up(invoices_gt: GroundTruth) -> None:
    for _, o in _occs(invoices_gt):
        if o.format not in ("pdf", "pdf_scanned"):
            continue
        e = o.expected
        line_sum = sum(Decimal(li["amount"]) for li in e["line_items"])
        sub = Decimal(e.get("subtotal", line_sum))
        consistent = line_sum == sub and sub + Decimal(e["tax"]) == Decimal(e["total"])
        bad = {"total_mismatch", "subtotal_mismatch"} & set(o.defects)
        assert consistent != bool(bad), o.occurrence_id


def test_text_pdfs_contain_their_fields(corpus: dict[str, bytes], invoices_gt: GroundTruth) -> None:
    text_pdfs = [o for _, o in _occs(invoices_gt, "pdf")][::7]  # every 7th, for speed
    assert len(text_pdfs) >= 15
    for o in text_pdfs:
        with pdfplumber.open(io.BytesIO(corpus[o.source])) as pdf:
            assert len(pdf.pages) == o.location["pages"]
            text = "\n".join(p.extract_text() or "" for p in pdf.pages)
        assert o.expected["invoice_number"] in text
        assert o.expected["vendor_name"] in text


def test_scanned_pdfs_have_no_text_layer(
    corpus: dict[str, bytes], invoices_gt: GroundTruth
) -> None:
    scans = [o for _, o in _occs(invoices_gt, "pdf_scanned")]
    for o in scans:
        with pdfplumber.open(io.BytesIO(corpus[o.source])) as pdf:
            assert all(not (p.extract_text() or "").strip() for p in pdf.pages), o.source
            assert pdf.pages[0].images


def test_exact_resends_are_byte_identical(
    corpus: dict[str, bytes], invoices_gt: GroundTruth
) -> None:
    by_id = {o.occurrence_id: o for _, o in _occs(invoices_gt)}
    exact = [o for o in by_id.values() if "exact_copy" in o.defects]
    assert len(exact) == CFG.n_resend
    for o in exact:
        assert o.duplicate_of
        assert corpus[o.source] == corpus[by_id[o.duplicate_of].source]


def test_legacy_ap_csv_is_messy_and_traceable(
    corpus: dict[str, bytes], invoices_gt: GroundTruth
) -> None:
    raw = corpus["invoices/ap_register.csv"]
    with pytest.raises(UnicodeDecodeError):
        raw.decode("utf-8")
    lines = raw.decode("latin-1").splitlines()
    assert lines[0] == "AP INVOICE REGISTER"
    assert lines[-1] == ",,,,,,,,"
    rows = list(csv.reader(lines))
    dates = [row[2] for row in rows[3:] if row[0]]
    assert any("/" in d for d in dates)
    assert any("-" in d for d in dates)
    assert any("," not in d and d[0].isalpha() for d in dates), "a date with no year"
    assert any(row[4].startswith("$") for row in rows[3:] if row[0])
    for _, o in _occs(invoices_gt, "csv"):
        if o.source != "invoices/ap_register.csv":
            continue
        row = rows[int(o.location["row"]) - 1]
        assert row[0] == o.expected["vendor_name"]
        assert row[1] == o.expected["invoice_number"]


def test_procurement_csv_has_aliases_gaps_and_duplicates(
    corpus: dict[str, bytes], invoices_gt: GroundTruth
) -> None:
    raw = corpus["invoices/procurement_export.csv"]
    assert raw.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(raw.decode("utf-8-sig").splitlines()))
    assert rows[0][1] == "supplier"
    assert any(row[2] == "" for row in rows[1:]), "missing invoice numbers"
    occs = [o for _, o in _occs(invoices_gt, "csv") if o.source.endswith("procurement_export.csv")]
    assert sum("duplicate" in o.defects for o in occs) == CFG.csv_b_dup_rows
    assert sum("vendor_alias" in o.defects for o in occs) > len(occs) // 2
    for o in occs:
        row = rows[int(o.location["row"]) - 1]
        assert row[1] == o.expected["vendor_name"]
        assert row[6] == o.expected["total"]


def test_aging_xlsx_layout(corpus: dict[str, bytes], invoices_gt: GroundTruth) -> None:
    wb = load_workbook(io.BytesIO(corpus["invoices/ap_aging.xlsx"]))
    assert wb.sheetnames == ["Notes", "Open Items"]
    ws = wb["Open Items"]
    assert ws["A1"].value == "Accounts Payable Aging Report"
    assert [c.value for c in ws[5]][:3] == ["Supplier", "Document", "Doc Date"]
    occs = [o for _, o in _occs(invoices_gt, "xlsx")]
    assert sum("stale_status" in o.defects for o in occs) == CFG.stale_aging_rows
    assert any("date_as_serial" in o.defects for o in occs)
    assert any("supplier_fill_down" in o.defects for o in occs)
    for o in occs:
        row = int(o.location["row"])
        assert ws.cell(row, 2).value == o.expected["invoice_number"]
        if "supplier_fill_down" in o.defects:
            assert ws.cell(row, 1).value is None


def test_lookalike_vendors_have_their_own_invoices(invoices_gt: GroundTruth) -> None:
    lookalikes = {e.entity_id for e in invoices_gt.entities if e.lookalike_of}
    assert len(lookalikes) == 3
    assert lookalikes & {r.entity_id for r in invoices_gt.records}


def test_currencies_present(invoices_gt: GroundTruth) -> None:
    assert {r.truth["currency"] for r in invoices_gt.records} == {"USD", "EUR", "CAD"}
