"""Each adapter, on small hand-built inputs."""

import io
import json
import random
from email.message import EmailMessage
from pathlib import Path

import docx
from openpyxl import Workbook
from reportlab.pdfgen import canvas

from data_gen.writers.scan import scan_pdf
from pipeline.ingest.adapters.eml import split_quoted
from pipeline.ingest.adapters.jsonish import flatten, parse_log_line
from pipeline.ingest.adapters.tabular import find_header
from pipeline.ingest.ingest import ingest_bytes, ingest_paths
from pipeline.ingest.models import location_covers


def test_csv_preamble_encoding_blank_rows() -> None:
    data = (
        "AP REGISTER\r\nExported today\r\nVendor,Inv Date,Amount\r\n"
        'Côté Mécanique Inc.,04/03/24,"$1,234.50"\r\nAcme,April 3,99\r\n,,\r\n,,\r\n'
    ).encode("latin-1")
    f = ingest_bytes("ap.csv", data)
    assert f.status == "ok"
    assert f.encoding == "cp1252"
    assert f.meta["preamble"] == ["AP REGISTER", "Exported today"]
    assert f.meta["blank_rows"] == 2
    assert [r.location for r in f.records] == [{"row": 4}, {"row": 5}]
    first, second = f.records
    assert first.raw["Vendor"] == "Côté Mécanique Inc."
    assert first.normalized["Inv Date"]["value"] == "2024-04-03"
    assert first.normalized["Amount"] == {
        "type": "amount",
        "value": "1234.50",
        "currency": "USD",
        "flags": ["dollar_sign_assumed_usd"],
    }
    assert second.normalized["Inv Date"]["partial"] == "--04-03"
    assert second.normalized["Amount"]["value"] == "99"


def test_csv_quoted_newline_keeps_physical_row_numbers() -> None:
    data = b'a,b\n1,"two\nlines"\n3,4\n'
    f = ingest_bytes("x.csv", data)
    assert [r.location["row"] for r in f.records] == [2, 4]


def test_find_header_skips_titles_and_spacers() -> None:
    rows: list[list[object]] = [
        ["Report"],
        [],
        ["Name", "Qty", "Note"],
        [],
        ["x", 1, None],
        ["y", 2, "z"],
    ]
    assert find_header(rows) == 2
    assert find_header([["just one cell"]]) is None


def _xlsx() -> bytes:
    wb = Workbook()
    notes = wb.active
    assert notes is not None
    notes.title = "Notes"
    notes["A1"] = "free text"
    ws = wb.create_sheet("Items")
    ws["A1"] = "Aging Report"
    ws.append([])
    ws.append(["Supplier", "Doc", "Doc Date", "Amount", "Days Past Due"])
    ws["G4"] = "stray note"
    ws.append(["Acme", "A-1", 45292, 10.5, 3])
    ws.append([None, "A-2", "2024/01/05", 2.25, 0])
    ws.append(["Total Acme", None, None, 12.75, None])
    ws.append(["Bolt", "B-1", None, 1, 0])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_xlsx_header_offset_fill_down_totals_serial_dates() -> None:
    f = ingest_bytes("aging.xlsx", _xlsx())
    assert f.status == "ok"
    kinds = [(r.kind, r.location) for r in f.records]
    assert kinds == [
        ("sheet_text", {"sheet": "Notes"}),
        ("stray_note", {"sheet": "Items", "row": 4}),
        ("row", {"sheet": "Items", "row": 5}),
        ("row", {"sheet": "Items", "row": 6}),
        ("total_row", {"sheet": "Items", "row": 7}),
        ("row", {"sheet": "Items", "row": 8}),
    ]
    a1, a2, bolt = f.records[2], f.records[3], f.records[5]
    assert a1.normalized["Doc Date"]["value"] == "2024-01-01"
    assert a1.normalized["Doc Date"]["flags"] == ["excel_serial_date"]
    assert "Days Past Due" not in a1.normalized
    assert a2.raw["Supplier"] == "Acme"
    assert a2.flags == ["filled_down:Supplier"]
    assert a2.normalized["Doc Date"]["value"] == "2024-01-05"
    assert bolt.flags == []  # fill-down resets after a total row
    json.dumps([r.raw for r in f.records])  # stored as JSONB: must be serializable
    assert f.meta["header_rows"] == {"Items": 3}


def _pdf(text: str) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, invariant=True)
    c.drawString(72, 720, text)
    c.showPage()
    c.drawString(72, 720, "page two")
    c.showPage()
    c.save()
    return buf.getvalue()


def test_pdf_text_and_scanned() -> None:
    text_pdf = _pdf("INVOICE 12345")
    f = ingest_bytes("a.pdf", text_pdf)
    assert f.format == "pdf"
    (rec,) = f.records
    assert rec.location == {"pages": 2}
    assert "--- page 1 of 2 ---" in rec.text and "INVOICE 12345" in rec.text
    assert rec.flags == ["multi_page"]

    scanned, _ = scan_pdf(text_pdf, random.Random(1))
    g = ingest_bytes("b.pdf", scanned)
    assert g.format == "pdf_scanned"
    assert g.records[0].text == ""
    assert g.records[0].flags == ["needs_vision", "no_text_layer"]


def _email(html_only: bool = False) -> bytes:
    msg = EmailMessage()
    msg["From"] = '"Dana Whitfield" <dana@northgate.example>'
    msg["To"] = "ops@ourco.example"
    msg["Subject"] = "Hi"
    msg["Date"] = "Tue, 21 Mar 2023 16:28:00 +0000"
    body = (
        "Hello,\n\nThanks.\n--\nDana Whitfield\n\n"
        "On Wed, 01 Mar 2023 09:10:00 +0000, Raj Patel <raj@bluewater.example> wrote:\n"
        "> Earlier note\n> Raj Patel\n"
    )
    if html_only:
        msg.set_content("<p>Hello</p><p>Dana</p>", subtype="html")
    else:
        msg.set_content(body)
        msg.add_attachment(b"BEGIN:VCARD\r\nFN:Dana\r\nEND:VCARD\r\n", maintype="text",
                           subtype="vcard", filename="dana.vcf")  # fmt: skip
        msg.add_attachment(b"%PDF-1.4", maintype="application", subtype="pdf", filename="q.pdf")
    return msg.as_bytes()


def test_eml_splits_body_quoted_and_attachments() -> None:
    f = ingest_bytes("m.eml", _email())
    parts = {r.location["part"]: r for r in f.records}
    assert set(parts) == {"body", "quoted", "attachment:dana.vcf", "attachment:q.pdf"}
    assert "dana@northgate.example" in parts["body"].text
    assert "Raj Patel" not in parts["body"].text
    assert parts["quoted"].text.startswith("On Wed, 01 Mar 2023")
    assert not any(line.startswith(">") for line in parts["quoted"].text.splitlines())
    assert "FN:Dana" in parts["attachment:dana.vcf"].text
    assert parts["attachment:q.pdf"].flags == ["binary_attachment"]
    assert parts["body"].normalized["Date"]["value"] == "2023-03-21"


def test_eml_html_only() -> None:
    f = ingest_bytes("m.eml", _email(html_only=True))
    body = f.records[0]
    assert body.flags == ["html_only"]
    assert "<p>" not in body.text and "Hello" in body.text


def test_split_quoted_with_bare_quote_lines() -> None:
    own, quoted = split_quoted("mine\n> theirs\n> more")
    assert own == "mine"
    assert quoted == "theirs\nmore"
    assert split_quoted("no quotes here") == ("no quotes here", None)


def test_docx_paragraphs_and_table_rows() -> None:
    d = docx.Document()
    d.add_heading("Contacts", level=1)
    d.add_paragraph("")
    d.add_paragraph("Primary contact: Dana Whitfield")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Name", "Email"
    row = t.add_row().cells
    row[0].text, row[1].text = "Dana Whitfield", "dana@northgate.example"
    buf = io.BytesIO()
    d.save(buf)
    f = ingest_bytes("c.docx", buf.getvalue())
    assert [(r.kind, r.location) for r in f.records] == [
        ("heading", {"paragraph": 0}),
        ("paragraph", {"paragraph": 2}),
        ("table_row", {"table": 0, "row": 1}),
    ]
    assert f.records[2].raw == {"Name": "Dana Whitfield", "Email": "dana@northgate.example"}


def test_text_and_markdown_cover_all_lines() -> None:
    f = ingest_bytes("n.md", b"# Title\n\n- Dana W.: 312-555-0143\n")
    (rec,) = f.records
    assert rec.location == {"line_start": 1, "line_end": 3}
    assert rec.flags == ["markdown"]
    assert location_covers(rec.location, {"line": 3})
    assert not location_covers(rec.location, {"line": 4})


def test_jsonl_shapes_and_malformed_lines() -> None:
    data = (
        b'{"name": "Dana", "contact": {"emailAddress": "d@x.example"}, "tel": null}\n'
        b"\n"
        b'{"event": "heartbeat", "ts": \n'
        b'{"data": {"submitted_at": "2024-06-12T07:19:00+00:00"}}\n'
    )
    f = ingest_bytes("w.jsonl", data)
    assert [(r.kind, r.location) for r in f.records] == [
        ("json_object", {"line": 1}),
        ("malformed_line", {"line": 3}),
        ("json_object", {"line": 4}),
    ]
    assert "contact.emailAddress: d@x.example" in f.records[0].text
    assert f.records[0].flags == ["has_nulls"]
    assert f.records[1].flags == ["malformed_json"]
    assert f.records[2].normalized["data.submitted_at"]["value"] == "2024-06-12"


def test_flatten() -> None:
    assert flatten({"a": {"b": 1}, "c": [{"d": 2}]}) == {"a.b": 1, "c.0.d": 2}


def test_log_formats() -> None:
    kv, flags = parse_log_line(
        '2024-01-04T17:35:00Z INFO  [webform] new lead name="Stephanie Laurent" '
        'company="Harborpoint Capital LLC" phone="+1 206 555 0194"'
    )
    assert flags == []
    assert kv["logger"] == "webform"
    assert kv["name"] == "Stephanie Laurent"
    assert kv["phone"] == "+1 206 555 0194"
    kv, _ = parse_log_line(
        '2024-01-06 09:34:00,147 INFO lead_capture: payload={"first": "Ann", "mail": "a@x.example"}'
    )
    assert kv["logger"] == "lead_capture"
    assert kv["payload.first"] == "Ann"
    kv, _ = parse_log_line("[2024-01-03 04:30:00] WARN crm: contact Elena <e@x.example> missing")
    assert kv["level"] == "WARN" and kv["logger"] == "crm"
    assert parse_log_line("garbage")[1] == ["unparsed_log_line"]


def test_json_array() -> None:
    f = ingest_bytes("a.json", b'[{"a": 1}, {"a": 2}]')
    assert [r.location for r in f.records] == [{"index": 0}, {"index": 1}]
    assert ingest_bytes("b.json", b"{nope").status == "error"


def test_unsupported_and_broken_files_do_not_stop_the_run() -> None:
    assert ingest_bytes("x.bin", b"\x00").status == "unsupported"
    broken = ingest_bytes("bad.xlsx", b"not a zip")
    assert broken.status == "error"
    assert broken.error and "BadZipFile" in broken.error
    assert broken.records == []


def test_ingest_paths_walks_dirs_and_excludes(tmp_path: Path) -> None:
    (tmp_path / "sub").mkdir()
    (tmp_path / "sub" / "a.txt").write_text("hello")
    (tmp_path / "ground_truth.json").write_text("{}")
    (tmp_path / ".DS_Store").write_bytes(b"\x00")
    files = ingest_paths([tmp_path])
    assert [f.path for f in files] == ["sub/a.txt"]
