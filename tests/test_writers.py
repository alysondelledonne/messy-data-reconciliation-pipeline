import io
import random
import zipfile
from email import message_from_bytes
from email.message import EmailMessage

import pdfplumber
import pytest
from openpyxl import Workbook, load_workbook

from data_gen.writers.files import csv_bytes, eml_bytes, normalize_zip, xlsx_bytes
from data_gen.writers.scan import scan_pdf


def _zip(data: bytes, when: tuple[int, int, int, int, int, int]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr(zipfile.ZipInfo("a.txt", date_time=when), data)
    return buf.getvalue()


def test_normalize_zip_ignores_write_time() -> None:
    a = _zip(b"hello", (2024, 1, 1, 9, 0, 0))
    b = _zip(b"hello", (2025, 6, 30, 17, 45, 10))
    assert a != b
    assert normalize_zip(a) == normalize_zip(b)


def test_xlsx_bytes_is_stable_and_readable() -> None:
    def make() -> bytes:
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws["A1"] = "x"
        return xlsx_bytes(wb)

    first = make()
    assert first == make()
    ws = load_workbook(io.BytesIO(first)).active
    assert ws is not None
    assert ws["A1"].value == "x"


def test_csv_bytes_latin1_preamble_and_blank_rows() -> None:
    data = csv_bytes(
        [["Vendor", "Amount"], ["Côté Mécanique Inc.", "$1,234.50"]],
        encoding="latin-1",
        preamble=["REPORT"],
        trailing_blank_rows=2,
    )
    with pytest.raises(UnicodeDecodeError):
        data.decode("utf-8")
    lines = data.decode("latin-1").split("\r\n")
    assert lines[0] == "REPORT"
    assert lines[2] == 'Côté Mécanique Inc.,"$1,234.50"'
    assert lines[3:5] == [",", ","]


def test_csv_bytes_bom() -> None:
    assert csv_bytes([["a"]], bom=True).startswith(b"\xef\xbb\xbf")


def test_eml_boundaries_are_fixed() -> None:
    def make() -> bytes:
        msg = EmailMessage()
        msg["From"] = "a@x.example"
        msg.set_content("hi")
        msg.add_attachment(b"BEGIN:VCARD", maintype="text", subtype="vcard", filename="a.vcf")
        return eml_bytes(msg, "k1")

    a = make()
    assert a == make()
    parsed = message_from_bytes(a)
    assert parsed.is_multipart()
    assert [p.get_filename() for p in parsed.walk() if p.get_filename()] == ["a.vcf"]


def _tiny_pdf() -> bytes:
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, invariant=True)
    c.drawString(72, 720, "INVOICE 12345 total $99.00")
    c.showPage()
    c.save()
    return buf.getvalue()


def test_scan_pdf_removes_text_layer_and_is_deterministic() -> None:
    src = _tiny_pdf()
    with pdfplumber.open(io.BytesIO(src)) as pdf:
        assert "12345" in (pdf.pages[0].extract_text() or "")
    a, params = scan_pdf(src, random.Random(9))
    b, _ = scan_pdf(src, random.Random(9))
    assert a == b
    with pdfplumber.open(io.BytesIO(a)) as pdf:
        assert (pdf.pages[0].extract_text() or "").strip() == ""
        assert pdf.pages[0].images  # it's an image now
    assert params["contrast"] < 1
