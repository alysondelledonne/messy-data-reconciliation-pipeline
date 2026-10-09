"""The contacts corpus holds the mess it claims to, and its ground truth matches the files."""

import csv
import io
import json
from collections import Counter
from email import message_from_bytes, policy
from email.message import EmailMessage
from email.utils import parseaddr

import docx

from data_gen.ground_truth import GroundTruth, Occurrence
from pipeline.domains.contacts.schema import Contact


def _occs(gt: GroundTruth, fmt: str) -> list[Occurrence]:
    return [o for r in gt.records for o in r.occurrences if o.format == fmt]


def test_truth_and_expected_values_fit_the_schema(contacts_gt: GroundTruth) -> None:
    for r in contacts_gt.records:
        Contact.model_validate(r.truth)
        for o in r.occurrences:
            Contact.model_validate({**r.truth, **o.expected})


def test_contact_details_are_reserved_fictional_ranges(contacts_gt: GroundTruth) -> None:
    for r in contacts_gt.records:
        assert r.truth["email"].endswith(".example")
        phone = r.truth["phone"]
        if phone.startswith("+1"):
            assert phone[5:9] == "5550" and phone[9] == "1", phone  # 555-0100..0199
        else:
            assert phone[:9] in {
                "+44207946", "+44161496", "+44117496", "+44113496",
            }  # fmt: skip


def test_namesakes_are_distinct_people(contacts_gt: GroundTruth) -> None:
    names = Counter((r.truth["first_name"], r.truth["last_name"]) for r in contacts_gt.records)
    assert sum(1 for n in names.values() if n > 1) == 3


def test_planted_mess_variety(contacts_gt: GroundTruth) -> None:
    defects = Counter(d for r in contacts_gt.records for o in r.occurrences for d in o.defects)
    for d in (
        "nickname", "company_alias", "stale_title", "stale_phone", "stale_company",
        "phone_transposed", "email_typo", "partial_name", "quoted_reply", "duplicate",
    ):  # fmt: skip
        assert defects[d] > 0, d
    assert len([e for e in contacts_gt.entities if e.lookalike_of]) == 4


def test_emails_parse_and_match_ground_truth(
    corpus: dict[str, bytes], contacts_gt: GroundTruth
) -> None:
    occs = _occs(contacts_gt, "eml")
    assert len({o.source for o in occs}) == 80
    for o in occs:
        msg = message_from_bytes(corpus[o.source], policy=policy.default)
        assert isinstance(msg, EmailMessage)
        part = str(o.location["part"])
        if part == "body":
            assert parseaddr(str(msg["From"]))[1] == o.expected["email"]
        elif part == "quoted":
            body = msg.get_body(preferencelist=("plain",))
            assert body is not None
            assert f"<{o.expected['email']}> wrote:" in body.get_content()
        else:
            fname = part.removeprefix("attachment:")
            atts = [a for a in msg.iter_attachments() if a.get_filename() == fname]
            assert atts and "BEGIN:VCARD" in atts[0].get_content()


def test_docx_locations_point_at_the_mention(
    corpus: dict[str, bytes], contacts_gt: GroundTruth
) -> None:
    docs: dict[str, docx.document.Document] = {}
    for o in _occs(contacts_gt, "docx"):
        d = docs.setdefault(o.source, docx.Document(io.BytesIO(corpus[o.source])))
        if "table" in o.location:
            row = d.tables[int(o.location["table"])].rows[int(o.location["row"])]
            text = " ".join(c.text for c in row.cells)
        else:
            text = d.paragraphs[int(o.location["paragraph"])].text
        assert o.expected["last_name"] in text
        if "email" in o.expected:
            assert o.expected["email"] in text


def test_notes_and_logs_locations(corpus: dict[str, bytes], contacts_gt: GroundTruth) -> None:
    for fmt in ("txt", "md", "log"):
        for o in _occs(contacts_gt, fmt):
            line = corpus[o.source].decode("utf-8").splitlines()[int(o.location["line"]) - 1]
            assert o.expected["first_name"] in line, (o.occurrence_id, line)


def test_jsonl_has_shapes_noise_and_repeats(
    corpus: dict[str, bytes], contacts_gt: GroundTruth
) -> None:
    lines = corpus["contacts/exports/webform_submissions.jsonl"].decode("utf-8").splitlines()
    bad = 0
    keysets = set()
    for line in lines:
        try:
            keysets.add(frozenset(json.loads(line)))
        except json.JSONDecodeError:
            bad += 1
    assert bad >= 1, "at least one malformed line"
    assert len(keysets) >= 3, "inconsistent shapes"
    occs = _occs(contacts_gt, "jsonl")
    assert sum("repeat_submission" in o.defects for o in occs) == 8
    for o in occs:
        json.loads(lines[int(o.location["line"]) - 1])  # every occurrence is a valid line


def test_crm_csv_duplicates(corpus: dict[str, bytes], contacts_gt: GroundTruth) -> None:
    rows = list(csv.reader(corpus["contacts/crm_export.csv"].decode("utf-8").splitlines()))
    assert rows[0][:3] == ["Contact ID", "First Name", "Last Name"]
    occs = _occs(contacts_gt, "csv")
    assert sum("duplicate" in o.defects for o in occs) == 10
    for o in occs:
        row = rows[int(o.location["row"]) - 1]
        assert row[1] == o.expected["first_name"]
