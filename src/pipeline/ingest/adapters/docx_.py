"""DOCX: one record per non-empty paragraph and one per table row (with its header)."""

import io

import docx

from pipeline.ingest.adapters.tabular import header_names, render
from pipeline.ingest.models import IngestedFile, RawRecord
from pipeline.ingest.normalize import normalize_row


def ingest_docx(data: bytes, f: IngestedFile) -> None:
    doc = docx.Document(io.BytesIO(data))
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if not text:
            continue
        style = para.style.name if para.style is not None else None
        kind = (
            "heading" if style and style.lower().startswith(("heading", "title")) else "paragraph"
        )
        f.records.append(RawRecord(kind, {"paragraph": i}, text, raw={"style": style}))
    for t, table in enumerate(doc.tables):
        rows = [[c.text.strip() for c in row.cells] for row in table.rows]
        if not rows:
            continue
        header = header_names(rows[0])
        for r, cells in enumerate(rows[1:], 1):
            if not any(cells):
                continue
            values = dict(zip(header, cells, strict=False))
            f.records.append(
                RawRecord("table_row", {"table": t, "row": r}, render(values), raw=values,
                          normalized=normalize_row(dict(values)))
            )  # fmt: skip
    f.meta["paragraphs"] = len(doc.paragraphs)
    f.meta["tables"] = len(doc.tables)
