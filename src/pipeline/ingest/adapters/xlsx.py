"""XLSX: one record per data row per sheet. Header not on row 1, serial dates,
total rows, fill-down for grouped reports, stray notes. A sheet with no table
becomes a single `sheet_text` record, so nothing is dropped."""

import io
from datetime import date, datetime, time
from typing import Any

from openpyxl import load_workbook

from pipeline.ingest.adapters.tabular import find_header, header_names, rows_to_records
from pipeline.ingest.models import IngestedFile, Location, RawRecord


def _jsonable(v: Any) -> Any:
    if isinstance(v, datetime | date | time):
        return v.isoformat()
    return v


def ingest_xlsx(data: bytes, f: IngestedFile) -> None:
    wb = load_workbook(io.BytesIO(data), data_only=True)
    f.meta["sheets"] = wb.sheetnames
    for ws in wb.worksheets:
        rows: list[tuple[int, list[Any]]] = []
        for cells in ws.iter_rows():
            rows.append((cells[0].row or 0, [c.value for c in cells]))
        if not rows:
            continue
        h = find_header([r for _, r in rows])
        if h is None:
            text = "\n".join(
                " ".join(str(v) for v in r if v not in (None, "")) for _, r in rows if any(r)
            )
            if text:
                f.records.append(RawRecord("sheet_text", {"sheet": ws.title}, text))
            continue
        header = header_names(rows[h][1])
        # Trim trailing col_<n> padding (the sheet's used range is often wider than the table).
        while header and header[-1].startswith("col_"):
            header.pop()
        f.meta.setdefault("header_rows", {})[ws.title] = rows[h][0]
        f.meta.setdefault("preamble", {})[ws.title] = [
            " ".join(str(v) for v in r if v not in (None, "")) for _, r in rows[:h] if any(r)
        ]
        data_rows: list[tuple[Location, list[Any]]] = [
            ({"sheet": ws.title, "row": n}, r) for n, r in rows[h + 1 :]
        ]
        records = rows_to_records(header, data_rows, fill_down_first_column=True)
        for rec in records:  # store raw values JSON-safe; normalized already has typed dates
            rec.raw = {k: [_jsonable(x) for x in v] if k == "_extra" else _jsonable(v)
                       for k, v in rec.raw.items()}  # fmt: skip
        f.records.extend(records)
